"""Tests for document parsers (PDF, PPTX, DOCX, and images)."""

import os
import tempfile
import pytest
import pymupdf
from pptx import Presentation
from docx import Document
from PIL import Image

from app.services.parsers import (
    parse_pdf,
    parse_pptx,
    parse_docx,
    parse_image,
    parse_document,
    is_tesseract_available,
)


@pytest.fixture
def temp_dir():
    """Create a temporary directory for test files."""
    with tempfile.TemporaryDirectory() as td:
        yield td


def test_parse_pdf_preserves_page_numbers(temp_dir):
    """Verify PyMuPDF extracts text per page with exact 1-indexed page numbers."""
    pdf_path = os.path.join(temp_dir, "sample.pdf")
    doc = pymupdf.open()

    page1 = doc.new_page()
    page1.insert_text((50, 50), "ClearDesk Page 1: Introduction to RAG Architecture")

    page2 = doc.new_page()
    page2.insert_text((50, 50), "ClearDesk Page 2: Vector search and citation chips")

    doc.save(pdf_path)
    doc.close()

    parsed = parse_pdf(pdf_path, file_id="file-123")
    assert parsed.page_count == 2
    assert len(parsed.pages) == 2

    assert parsed.pages[0].page_or_slide == 1
    assert "Page 1: Introduction to RAG Architecture" in parsed.pages[0].text
    assert parsed.pages[0].label == "Page 1"

    assert parsed.pages[1].page_or_slide == 2
    assert "Page 2: Vector search and citation chips" in parsed.pages[1].text
    assert parsed.pages[1].label == "Page 2"


def test_parse_pptx_preserves_slide_numbers_and_notes(temp_dir):
    """Verify python-pptx extracts slide text and speaker notes with exact slide numbers."""
    pptx_path = os.path.join(temp_dir, "sample.pptx")
    prs = Presentation()

    # Slide 1
    blank_layout = prs.slide_layouts[6]
    slide1 = prs.slides.add_slide(blank_layout)
    txBox = slide1.shapes.add_textbox(0, 0, 100, 100)
    txBox.text_frame.text = "Slide 1: Gradient Descent Fundamentals"

    # Slide 2 with speaker notes
    slide2 = prs.slides.add_slide(blank_layout)
    txBox2 = slide2.shapes.add_textbox(0, 0, 100, 100)
    txBox2.text_frame.text = "Slide 2: Convex Optimization"
    notes_slide = slide2.notes_slide
    notes_slide.notes_text_frame.text = "Remember to emphasize learning rate selection"

    prs.save(pptx_path)

    parsed = parse_pptx(pptx_path, file_id="file-pptx")
    assert parsed.page_count == 2
    assert parsed.pages[0].page_or_slide == 1
    assert "Slide 1: Gradient Descent Fundamentals" in parsed.pages[0].text
    assert parsed.pages[0].label == "Slide 1"

    assert parsed.pages[1].page_or_slide == 2
    assert "Slide 2: Convex Optimization" in parsed.pages[1].text
    assert "Speaker Notes: Remember to emphasize learning rate selection" in parsed.pages[1].text
    assert parsed.pages[1].label == "Slide 2"


def test_parse_docx_groups_by_sections(temp_dir):
    """Verify python-docx parses content into identifiable logical sections without claiming false page numbers."""
    docx_path = os.path.join(temp_dir, "sample.docx")
    doc = Document()
    doc.add_heading("Section 1: Project Overview", level=1)
    doc.add_paragraph("ClearDesk is an AI-powered document workspace.")
    doc.add_heading("Section 2: Anti-Hallucination Guardrails", level=1)
    doc.add_paragraph("Answers must strictly be grounded in source chunks.")
    doc.save(docx_path)

    parsed = parse_docx(docx_path, file_id="file-docx")
    assert parsed.page_count >= 2
    assert parsed.pages[0].page_or_slide == 1
    assert "Section 1" in parsed.pages[0].label
    assert "Project Overview" in parsed.pages[0].text

    assert parsed.pages[1].page_or_slide == 2
    assert "Section 2" in parsed.pages[1].label
    assert "Anti-Hallucination Guardrails" in parsed.pages[1].text


def test_corrupt_file_raises_value_error(temp_dir):
    """Verify corrupt files fail gracefully with clean error messages."""
    bad_pdf = os.path.join(temp_dir, "bad.pdf")
    with open(bad_pdf, "wb") as f:
        f.write(b"NOT_A_VALID_PDF_HEADER_DATA")

    with pytest.raises(ValueError, match="Corrupt or invalid PDF"):
        parse_pdf(bad_pdf, file_id="corrupt-1")


def test_unsupported_file_type_dispatcher():
    """Verify dispatcher rejects unsupported extensions."""
    with pytest.raises(ValueError, match="Unsupported file type"):
        parse_document("test.exe", "file-id", "exe")
